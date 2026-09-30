"""
OCR / Document extraction agent.

Supports:
    - Images: .jpg, .jpeg, .png, .bmp, .tiff
    - PDF: .pdf
    - Word: .docx, .doc

For images/PDFs:
    Uses Tesseract OCR.

For Word:
    Extracts text directly from the document.

Usage:
    python ocr_test.py sample_docs/invoice1.jpg
    python ocr_test.py sample_docs/invoice1.pdf
    python ocr_test.py sample_docs/invoice1.docx
"""

import sys
import os
import pytesseract

from PIL import Image, ImageOps, ImageFilter
from docx import Document

# PDF support
from pdf2image import convert_from_path


# If Tesseract isn't on your PATH, uncomment:
# pytesseract.pytesseract.tesseract_cmd = (
#     r"C:\Program Files\Tesseract-OCR\tesseract.exe"
# )

UPSCALE_FACTOR = 3

SUPPORTED_IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tiff",
    ".tif",
}

SUPPORTED_EXTENSIONS = SUPPORTED_IMAGE_EXTENSIONS | {
    ".pdf",
    ".doc",
    ".docx",
}


def preprocess_image(img: Image.Image) -> Image.Image:
    """
    Grayscale + upscale + autocontrast + sharpen.
    Helps OCR read smaller or lower-quality text.
    """

    processed = img.convert("L")

    processed = processed.resize(
        (
            processed.width * UPSCALE_FACTOR,
            processed.height * UPSCALE_FACTOR,
        ),
        Image.LANCZOS,
    )

    processed = ImageOps.autocontrast(processed)
    processed = processed.filter(ImageFilter.SHARPEN)

    return processed


def extract_text_from_image(image_path: str, raw: bool = False) -> str:
    """
    Extract text from an image using Tesseract OCR.
    """

    img = Image.open(image_path)

    if not raw:
        img = preprocess_image(img)

    return pytesseract.image_to_string(img)


def extract_text_from_pdf(pdf_path: str, raw: bool = False) -> str:
    """
    Convert each PDF page to an image and run OCR on it.

    This works for scanned/image-based PDFs as well as PDFs
    where the text layer is difficult to extract.
    """

    print("Converting PDF pages to images...")

    pages = convert_from_path(pdf_path, dpi=300)

    all_text = []

    for page_number, page in enumerate(pages, start=1):

        print(f"Processing PDF page {page_number}/{len(pages)}...")

        if not raw:
            page = preprocess_image(page)

        page_text = pytesseract.image_to_string(page)

        all_text.append(
            f"\n--- PDF PAGE {page_number} ---\n"
            f"{page_text}"
        )

    return "\n".join(all_text)


def extract_text_from_docx(docx_path: str) -> str:
    """
    Extract text from a .docx Word document.

    Reads paragraphs and table contents.
    """

    document = Document(docx_path)

    text_parts = []

    # Paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_data = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                row_data.append(cell_text)

            if any(row_data):
                text_parts.append(" | ".join(row_data))

    return "\n".join(text_parts)


def extract_text_from_doc(doc_path: str) -> str:
    """
    .doc (old Microsoft Word format) support.

    This requires LibreOffice to be installed.
    Converts .doc -> .docx temporarily and then extracts text.
    """

    import subprocess
    import tempfile

    temp_dir = tempfile.mkdtemp()

    try:

        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                temp_dir,
                doc_path,
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

        converted_file = os.path.join(
            temp_dir,
            os.path.splitext(os.path.basename(doc_path))[0] + ".docx",
        )

        if not os.path.exists(converted_file):
            raise RuntimeError(
                "LibreOffice could not convert the .doc file."
            )

        return extract_text_from_docx(converted_file)

    except FileNotFoundError:
        raise RuntimeError(
            "LibreOffice is required to process .doc files. "
            "Install LibreOffice or use .docx instead."
        )


def extract_text(file_path: str, raw: bool = False) -> str:
    """
    Automatically detects the file type and extracts text.
    """

    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(file_path)[1].lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}\n"
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    # IMAGE
    if extension in SUPPORTED_IMAGE_EXTENSIONS:
        return extract_text_from_image(
            file_path,
            raw=raw
        )

    # PDF
    elif extension == ".pdf":
        return extract_text_from_pdf(
            file_path,
            raw=raw
        )

    # DOCX
    elif extension == ".docx":
        return extract_text_from_docx(file_path)

    # DOC
    elif extension == ".doc":
        return extract_text_from_doc(file_path)

    return ""


if __name__ == "__main__":

    if len(sys.argv) < 2:
        print(
            "Usage: python ocr_test.py "
            "<file_path> [--raw]"
        )
        sys.exit(1)

    path = sys.argv[1]
    raw = "--raw" in sys.argv

    try:

        result = extract_text(
            path,
            raw=raw
        )

        print(
            f"\n----- EXTRACTED TEXT "
            f"FROM {os.path.basename(path)} -----"
        )

        print(result)

        print("---------------------------")

    except Exception as e:

        print(f"ERROR: {e}")
        sys.exit(1)